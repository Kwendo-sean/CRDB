import tempfile
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.importers import import_activity_workbook, import_participant_workbook, import_quiz_workbook, import_session_workbook
from core.models import Activity, Experience, Participant, Quiz, Session
from core.workbooks import TEMPLATES, quiz_template, template_bytes


def write(workbook):
    path = Path(tempfile.mkdtemp()) / "template.xlsx"
    workbook.save(path)
    return str(path)


class TemplateShapeTests(TestCase):
    def test_every_template_carries_instructions_and_a_data_sheet(self):
        for key, (filename, builder) in TEMPLATES.items():
            with self.subTest(key=key):
                workbook = load_workbook(write(builder()))
                self.assertEqual(workbook.sheetnames[0], "INSTRUCTIONS", key)
                self.assertEqual(len(workbook.sheetnames), 2, key)
                data = workbook[workbook.sheetnames[1]]
                self.assertGreaterEqual(data.max_row, 3, "template needs a header and worked examples")
                self.assertTrue(filename.endswith(".xlsx"))

    def test_template_bytes_returns_a_real_xlsx(self):
        filename, payload = template_bytes("participants")
        self.assertEqual(filename, "participants-template.xlsx")
        self.assertTrue(payload.startswith(b"PK"))


class TemplateRoundTripTests(TestCase):
    """The worked example rows shipped in each template must actually import."""

    def test_session_template_examples_import(self):
        summary = import_session_workbook(write(TEMPLATES["programme-sessions"][1]()))
        self.assertEqual(summary["sessions"], 2)
        self.assertEqual(Session.objects.count(), 2)

    def test_session_reimport_updates_rather_than_duplicates(self):
        path = write(TEMPLATES["programme-sessions"][1]())
        import_session_workbook(path)
        summary = import_session_workbook(path)
        self.assertEqual(summary["updated"], 2)
        self.assertEqual(Session.objects.count(), 2)

    def test_activity_template_examples_import(self):
        summary = import_activity_workbook(write(TEMPLATES["activities"][1]()))
        self.assertEqual(summary["activities"], 2)
        self.assertTrue(Activity.objects.filter(code="fraud-hunt", points=200, repeatable=False).exists())

    def test_participant_template_examples_import(self):
        summary = import_participant_workbook(write(TEMPLATES["participants"][1]()))
        self.assertEqual((summary["created"], summary["updated"]), (2, 0))
        person = Participant.objects.get(staff_id="CRDB0142")
        self.assertEqual(person.department, "Retail Banking")
        self.assertTrue(person.pass_number)

    def test_participant_reimport_updates_and_keeps_the_original_pass(self):
        path = write(TEMPLATES["participants"][1]())
        import_participant_workbook(path)
        original = Participant.objects.get(staff_id="CRDB0142")
        summary = import_participant_workbook(path)
        original.refresh_from_db()
        self.assertEqual((summary["created"], summary["updated"]), (0, 2))
        self.assertEqual(Participant.objects.count(), 2)
        self.assertEqual(Participant.objects.get(staff_id="CRDB0142").pass_number, original.pass_number)

    def test_quiz_template_examples_are_rejected_until_all_ten_exist(self):
        with self.assertRaises(ValidationError) as caught:
            import_quiz_workbook(write(quiz_template()))
        self.assertIn("exactly 10", " ".join(caught.exception.messages))

    def test_quiz_template_imports_once_it_holds_ten_questions(self):
        workbook = quiz_template()
        sheet = workbook["QUESTIONS"]
        for number in range(3, 11):
            sheet.append([1, "Systems Online", number, f"Question {number}", "A", "B", "C", "D", 1, 20, 1000, "Because."])
        summary = import_quiz_workbook(write(workbook))
        self.assertEqual(summary, {"quizzes": 1, "questions": 10})
        quiz = Quiz.objects.get()
        self.assertEqual(quiz.questions.count(), 10)
        self.assertEqual(quiz.questions.first().choices.filter(is_correct=True).count(), 1)


class ImporterRejectionTests(TestCase):
    def test_duplicate_staff_id_in_one_file_is_reported_by_row(self):
        workbook = TEMPLATES["participants"][1]()
        sheet = workbook["PARTICIPANTS"]
        sheet.append(["CRDB0142", "Grace Twin", "other@crdbbank.co.tz", "", "", ""])
        with self.assertRaises(ValidationError) as caught:
            import_participant_workbook(write(workbook))
        self.assertIn("Row 4", " ".join(caught.exception.messages))
        self.assertEqual(Participant.objects.count(), 0, "a rejected file must import nothing")

    def test_missing_columns_name_the_template(self):
        workbook = TEMPLATES["activities"][1]()
        workbook["ACTIVITIES"].delete_cols(5)
        with self.assertRaises(ValidationError) as caught:
            import_activity_workbook(write(workbook))
        self.assertIn("points", caught.exception.messages[0])

    def test_activity_row_with_a_bad_day_blocks_the_whole_file(self):
        workbook = TEMPLATES["activities"][1]()
        workbook["ACTIVITIES"].append(["late-stand", "Late Stand", "", 9, 100, "NO"])
        with self.assertRaises(ValidationError) as caught:
            import_activity_workbook(write(workbook))
        self.assertIn("day must be 1-4", " ".join(caught.exception.messages))
        self.assertEqual(Activity.objects.count(), 0)


class ControlRoomImportTests(TestCase):
    def setUp(self):
        get_user_model().objects.create_user("ops", password="control-room-pass", is_staff=True)
        self.client.login(username="ops", password="control-room-pass")

    def test_template_downloads_are_staff_only(self):
        self.client.logout()
        response = self.client.get(reverse("control-template-download", args=["participants"]))
        self.assertEqual(response.status_code, 302)

    def test_template_download_serves_a_spreadsheet(self):
        response = self.client.get(reverse("control-template-download", args=["quiz-questions"]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("quiz-questions-template.xlsx", response["Content-Disposition"])

    def test_unknown_template_key_redirects_instead_of_erroring(self):
        response = self.client.get(reverse("control-template-download", args=["nope"]))
        self.assertEqual(response.status_code, 302)

    def test_participant_upload_creates_people(self):
        with open(write(TEMPLATES["participants"][1]()), "rb") as handle:
            response = self.client.post(reverse("control-participants-import"), {"workbook": handle}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Participant.objects.count(), 2)

    def test_activity_upload_also_creates_the_participant_experience(self):
        with open(write(TEMPLATES["activities"][1]()), "rb") as handle:
            self.client.post(reverse("control-activities-import"), {"workbook": handle}, follow=True)
        self.assertEqual(Activity.objects.count(), 2)
        self.assertEqual(Experience.objects.filter(type=Experience.Type.EXPO_ACTIVITY).count(), 2)

    def test_rejected_upload_reports_the_row_and_saves_nothing(self):
        workbook = TEMPLATES["activities"][1]()
        workbook["ACTIVITIES"].append(["bad-row", "Bad Row", "", 9, 100, "NO"])
        with open(write(workbook), "rb") as handle:
            response = self.client.post(reverse("control-activities-import"), {"workbook": handle}, follow=True)
        self.assertContains(response, "IMPORT REJECTED")
        self.assertEqual(Activity.objects.count(), 0)

    def test_wrong_file_type_is_refused(self):
        path = Path(tempfile.mkdtemp()) / "roster.txt"
        path.write_text("not a spreadsheet")
        with open(path, "rb") as handle:
            response = self.client.post(reverse("control-participants-import"), {"workbook": handle}, follow=True)
        self.assertContains(response, "IMPORT REJECTED")
        self.assertEqual(Participant.objects.count(), 0)

    def test_missing_file_does_not_error(self):
        response = self.client.post(reverse("control-sessions-import"), {}, follow=True)
        self.assertContains(response, "NO SPREADSHEET WAS ATTACHED")


class SampleQuizTests(TestCase):
    """The shipped worked example must always be importable as-is."""

    def test_the_sample_imports_without_edits(self):
        summary = import_quiz_workbook(write(TEMPLATES["sample-quiz"][1]()))
        self.assertEqual(summary, {"quizzes": 1, "questions": 10})

    def test_every_sample_question_is_publishable(self):
        import_quiz_workbook(write(TEMPLATES["sample-quiz"][1]()))
        quiz = Quiz.objects.get()
        self.assertEqual(quiz.questions.count(), 10)
        self.assertEqual(sorted(quiz.questions.values_list("number", flat=True)), list(range(1, 11)))
        for question in quiz.questions.prefetch_related("choices"):
            choices = list(question.choices.all())
            self.assertEqual(len(choices), 4, f"Q{question.number}")
            self.assertEqual(sum(choice.is_correct for choice in choices), 1, f"Q{question.number}")
            self.assertTrue(question.host_note.strip(), f"Q{question.number} needs a host note")
            self.assertIn(question.time_limit_seconds, (20, 30))

    def test_the_sample_passes_the_publication_gate(self):
        import_quiz_workbook(write(TEMPLATES["sample-quiz"][1]()))
        quiz = Quiz.objects.get()
        get_user_model().objects.create_user("ops", password="control-room-pass", is_staff=True)
        self.client.login(username="ops", password="control-room-pass")
        self.client.post(reverse("control-quiz-publish", args=[quiz.pk]))
        quiz.refresh_from_db()
        self.assertEqual(quiz.state, Quiz.State.SCHEDULED)

    def test_the_sample_is_downloadable_from_the_control_room(self):
        get_user_model().objects.create_user("ops2", password="control-room-pass", is_staff=True)
        self.client.login(username="ops2", password="control-room-pass")
        response = self.client.get(reverse("control-template-download", args=["sample-quiz"]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("sample-quiz-day1", response["Content-Disposition"])


class ExpoGamesSheetTests(TestCase):
    """The shipped expo-games sheet must load the real stand line-up as-is."""

    def test_it_imports_the_three_games(self):
        from core.workbooks import expo_games
        summary = import_activity_workbook(write(expo_games()))
        self.assertEqual(summary["created"], 3)
        self.assertEqual(
            sorted(Activity.objects.values_list("code", flat=True)),
            ["arduino-assembly", "robot-assembly", "vr-games"],
        )

    def test_the_games_carry_their_points(self):
        from core.workbooks import expo_games
        import_activity_workbook(write(expo_games()))
        self.assertEqual(Activity.objects.get(code="vr-games").points, 200)
        self.assertEqual(Activity.objects.get(code="robot-assembly").points, 250)
        self.assertEqual(Activity.objects.get(code="arduino-assembly").points, 250)

    def test_reimporting_does_not_duplicate_or_disturb_awarded_points(self):
        from core.workbooks import expo_games
        path = write(expo_games())
        import_activity_workbook(path)
        person = Participant.objects.create(full_name="Grace Mushi", staff_id="EG1")
        from core.services import award_activity
        award_activity(participant=person, activity=Activity.objects.get(code="vr-games"))
        summary = import_activity_workbook(path)
        self.assertEqual((summary["created"], summary["updated"]), (0, 3))
        self.assertEqual(Activity.objects.count(), 3)
        self.assertEqual(person.activity_completions.count(), 1)
