import csv
from collections import defaultdict
from datetime import date,datetime,time
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from .models import Activity, AnswerChoice, Participant, Question, Quiz, QuizResponse, Session

QUESTIONS_PER_CHALLENGE = 10

REQUIRED=("day","challenge_name","question_number","question_text","answer_1","answer_2","answer_3","answer_4","correct_answer_number","time_limit","points","host_note")
def _clean_header(value): return str(value or "").strip().lower().replace(" ","_")
def _csv_rows(source):
    with open(source, newline="", encoding="utf-8-sig") as handle:
        rows=list(csv.reader(handle))
    if not rows: raise ValidationError("The file is empty.")
    return rows[0], rows[1:]

def _sheet_rows(source, sheet_name=None):
    workbook=load_workbook(source,data_only=True,read_only=True)
    try:
        sheet=workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
        rows=sheet.iter_rows(values_only=True)
        try: header=next(rows)
        except StopIteration: raise ValidationError("The sheet has no header row.")
        return list(header), list(rows)
    finally:
        # read_only mode holds the file open; the caller deletes the upload straight after.
        workbook.close()

def _tabular_rows(source, sheet_name=None):
    """Read the header and data rows of a CSV or XLSX file."""
    if str(source).lower().endswith(".csv"): return _csv_rows(source)
    return _sheet_rows(source, sheet_name)

def _blank(values):
    return not any(value is not None and str(value).strip() for value in values)

def _read_table(source, required, sheet_name=None):
    raw_headers, raw_rows = _tabular_rows(source, sheet_name)
    headers=[_clean_header(value) for value in raw_headers]
    missing=[column for column in required if column not in headers]
    if missing: raise ValidationError(f"Missing columns: {', '.join(missing)}. Use the downloadable template.")
    return headers, raw_rows

def _quiz_rows(source):
    return _tabular_rows(source, "QUESTIONS")

def parse_quiz_workbook(source):
    headers, raw_rows = _read_table(source, REQUIRED, sheet_name="QUESTIONS")
    groups=defaultdict(list); errors=[]
    for row_number,values in enumerate(raw_rows,start=2):
        if _blank(values): continue
        data=dict(zip(headers,values))
        try:
            day=int(data["day"]); number=int(data["question_number"]); correct=int(data["correct_answer_number"]); timer=int(data["time_limit"]); points=int(data["points"])
            if day not in range(1,5): raise ValueError("day must be 1-4")
            if correct not in range(1,5): raise ValueError("correct answer must be 1-4")
            if not 5<=timer<=3600: raise ValueError("time_limit must be between 5 and 3600 seconds")
            if points<=0: raise ValueError("points must be positive")
            required_text=[data["challenge_name"],data["question_text"],*[data[f"answer_{i}"] for i in range(1,5)],data["host_note"]]
            if any(v is None or not str(v).strip() for v in required_text): raise ValueError("question, four answers and host note are required")
            groups[(day,str(data["challenge_name"]).strip())].append((number,data,correct,timer,points))
        except (TypeError,ValueError) as exc: errors.append(f"Row {row_number}: {exc}")
    for (day,title),rows in groups.items():
        numbers=[row[0] for row in rows]
        if len(numbers)!=len(set(numbers)):
            duplicates=sorted({number for number in numbers if numbers.count(number)>1})
            errors.append(f"Day {day} / {title}: question_number repeated ({', '.join(str(n) for n in duplicates)})")
        elif len(rows)!=QUESTIONS_PER_CHALLENGE or set(numbers)!=set(range(1,QUESTIONS_PER_CHALLENGE+1)):
            errors.append(f"Day {day} / {title}: found {len(rows)} questions numbered {sorted(numbers)}; an official challenge needs exactly {QUESTIONS_PER_CHALLENGE}, numbered 1-{QUESTIONS_PER_CHALLENGE}")
    if not groups: errors.append("No question rows were found. Fill in the QUESTIONS sheet.")
    if errors: raise ValidationError(errors)
    return groups

def select_quiz_questions(groups, quiz):
    if len(groups)==1: return next(iter(groups.values()))
    exact=groups.get((quiz.day,quiz.title))
    if exact: return exact
    day_matches=[rows for (day,_),rows in groups.items() if day==quiz.day]
    if len(day_matches)==1: return day_matches[0]
    raise ValidationError("The workbook must contain one challenge, or one unambiguous challenge matching this quiz.")

@transaction.atomic
def import_questions_into_quiz(quiz, source):
    rows=select_quiz_questions(parse_quiz_workbook(source),quiz)
    if QuizResponse.objects.filter(question__quiz=quiz).exists():
        raise ValidationError("This round already holds participant answers, so its questions cannot be replaced. Create a new round instead.")
    quiz.questions.all().delete()
    for number,data,correct,timer,points in sorted(rows):
        question=Question.objects.create(quiz=quiz,number=number,text=str(data["question_text"]).strip(),time_limit_seconds=timer,base_points=points,host_note=str(data["host_note"]).strip())
        AnswerChoice.objects.bulk_create([AnswerChoice(question=question,number=i,text=str(data[f"answer_{i}"]).strip(),is_correct=i==correct) for i in range(1,5)])
    return {"questions":len(rows)}

@transaction.atomic
def import_quiz_workbook(source):
    groups=parse_quiz_workbook(source)
    played=Quiz.objects.filter(sequence__in=[day for day,_ in groups],attempts__isnull=False).distinct()
    if played.exists():
        raise ValidationError("Rounds " + ", ".join(quiz.title for quiz in played) + " already have attempts and will not be replaced.")
    Quiz.objects.filter(sequence__in=[day for day,_ in groups]).delete()
    count=0
    for (day,title),rows in sorted(groups.items()):
        quiz=Quiz.objects.create(day=day,title=str(title),sequence=day,prize_eligible=True,speed_bonus_enabled=False)
        for number,data,correct,timer,points in sorted(rows):
            question=Question.objects.create(quiz=quiz,number=number,text=str(data["question_text"]).strip(),time_limit_seconds=timer,base_points=points,host_note=str(data["host_note"]).strip())
            AnswerChoice.objects.bulk_create([AnswerChoice(question=question,number=i,text=str(data[f"answer_{i}"]).strip(),is_correct=i==correct) for i in range(1,5)])
            count+=1
    return {"quizzes":len(groups),"questions":count}

SESSION_REQUIRED=("source_key","day","date","start_time","end_time","title","facilitator","format","audience","location")
def _parse_date(value):
    if isinstance(value,datetime): return value.date()
    if isinstance(value,date): return value
    return datetime.strptime(str(value).strip(),"%Y-%m-%d").date()
def _parse_time(value):
    if isinstance(value,datetime): return value.time()
    if isinstance(value,time): return value
    return datetime.strptime(str(value).strip(),"%H:%M").time()
@transaction.atomic
def import_session_workbook(source):
    headers,raw_rows=_read_table(source,SESSION_REQUIRED,sheet_name="SESSIONS")
    parsed=[]; errors=[]; seen=set()
    formats={"FACE-TO-FACE":"face","FACE TO FACE":"face","VIRTUAL":"virtual","HYBRID":"hybrid"}
    for row_number,values in enumerate(raw_rows,start=2):
        if _blank(values): continue
        data=dict(zip(headers,values))
        try:
            source_key=str(data["source_key"] or "").strip(); title=str(data["title"] or "").strip(); facilitator=str(data["facilitator"] or "").strip(); day=int(data["day"])
            if not source_key or not title: raise ValueError("source_key and title are required")
            if day not in range(1,5): raise ValueError("day must be 1-4")
            event_date=_parse_date(data["date"]); starts=timezone.make_aware(datetime.combine(event_date,_parse_time(data["start_time"]))); ends=timezone.make_aware(datetime.combine(event_date,_parse_time(data["end_time"])))
            if ends<=starts: raise ValueError("end_time must be after start_time")
            format_value=formats.get(str(data["format"] or "").strip().upper())
            if not format_value: raise ValueError("format must be FACE-TO-FACE, VIRTUAL or HYBRID")
            if source_key in seen: raise ValueError(f"source_key '{source_key}' is used more than once in this file")
            seen.add(source_key)
            parsed.append({"source_key":source_key,"day":day,"title":title,"facilitator":facilitator,"format":format_value,"audience":str(data["audience"] or "").strip(),"location":str(data["location"] or "").strip(),"starts_at":starts,"ends_at":ends})
        except (TypeError,ValueError) as exc: errors.append(f"Row {row_number}: {exc}")
    if not parsed and not errors: errors.append("No session rows were found. Fill in the SESSIONS sheet.")
    if errors: raise ValidationError(errors)
    created=updated=0
    for values in parsed:
        _,was_created=Session.objects.update_or_create(source_key=values.pop("source_key"),defaults=values)
        created+=was_created; updated+=not was_created
    return {"sessions":len(parsed),"created":created,"updated":updated}


PARTICIPANT_REQUIRED=("staff_id","full_name","email","department","job_role","phone")

@transaction.atomic
def import_participant_workbook(source):
    """Bulk-load the registered roster.

    Rows are matched on staff ID first, then email, so re-importing a corrected
    file updates people instead of duplicating them. Existing passes, QR tokens
    and scores are never touched.
    """
    headers,raw_rows=_read_table(source,PARTICIPANT_REQUIRED,sheet_name="PARTICIPANTS")
    created=updated=0; errors=[]; seen_staff=set(); seen_email=set()
    for row_number,values in enumerate(raw_rows,start=2):
        if _blank(values): continue
        data=dict(zip(headers,values))
        try:
            staff=str(data.get("staff_id") or "").strip().upper() or None
            email=str(data.get("email") or "").strip().lower() or None
            name=str(data.get("full_name") or "").strip()
            if not name: raise ValueError("full_name is required")
            if not staff and not email: raise ValueError("staff_id or email is required")
            if staff and staff in seen_staff: raise ValueError(f"staff_id '{staff}' appears more than once in this file")
            if email and email in seen_email: raise ValueError(f"email '{email}' appears more than once in this file")
            by_staff=Participant.objects.filter(staff_id=staff).first() if staff else None
            by_email=Participant.objects.filter(email=email).first() if email else None
            if by_staff and by_email and by_staff.pk!=by_email.pk:
                raise ValueError("staff_id and email already belong to two different participants")
            participant=by_staff or by_email
            fields={"full_name":name,"staff_id":staff,"email":email,
                    "department":str(data.get("department") or "").strip(),
                    "job_role":str(data.get("job_role") or "").strip(),
                    "phone":str(data.get("phone") or "").strip()}
            if participant:
                for key,value in fields.items():
                    if value: setattr(participant,key,value)
                participant.save(); updated+=1
            else:
                Participant.objects.create(**fields); created+=1
            if staff: seen_staff.add(staff)
            if email: seen_email.add(email)
        except (TypeError,ValueError,ValidationError) as exc:
            message=" ".join(exc.messages) if isinstance(exc,ValidationError) else str(exc)
            errors.append(f"Row {row_number}: {message}")
    if not created and not updated and not errors:
        errors.append("No participant rows were found. Fill in the PARTICIPANTS sheet.")
    if errors: raise ValidationError(errors)
    return {"participants":created+updated,"created":created,"updated":updated}


ACTIVITY_REQUIRED=("code","name","description","day","points","repeatable")

@transaction.atomic
def import_activity_workbook(source):
    """Load interactive games, expo stands and AIoT demos keyed on their code."""
    headers,raw_rows=_read_table(source,ACTIVITY_REQUIRED,sheet_name="ACTIVITIES")
    created=updated=0; errors=[]; seen=set()
    truthy={"YES":True,"Y":True,"TRUE":True,"1":True,"NO":False,"N":False,"FALSE":False,"0":False,"":False}
    for row_number,values in enumerate(raw_rows,start=2):
        if _blank(values): continue
        data=dict(zip(headers,values))
        try:
            code=str(data.get("code") or "").strip().lower()
            name=str(data.get("name") or "").strip()
            if not code or not name: raise ValueError("code and name are required")
            if code in seen: raise ValueError(f"code '{code}' appears more than once in this file")
            seen.add(code)
            day=int(data["day"])
            if day not in range(1,5): raise ValueError("day must be 1-4")
            points=int(data["points"])
            if points<0: raise ValueError("points cannot be negative")
            repeatable=truthy.get(str(data.get("repeatable") or "").strip().upper())
            if repeatable is None: raise ValueError("repeatable must be YES or NO")
            _,was_created=Activity.objects.update_or_create(code=code,defaults={"name":name,"description":str(data.get("description") or "").strip(),"day":day,"points":points,"repeatable":repeatable,"is_active":True})
            created+=was_created; updated+=not was_created
        except (KeyError,TypeError,ValueError) as exc:
            errors.append(f"Row {row_number}: {exc}")
    if not created and not updated and not errors:
        errors.append("No activity rows were found. Fill in the ACTIVITIES sheet.")
    if errors: raise ValidationError(errors)
    return {"activities":created+updated,"created":created,"updated":updated}
